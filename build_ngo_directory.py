#!/usr/bin/env python3
"""Build the LIKE NGO Directory department pages from the workbook snapshots.

Reads:
  ~/workspace/ngo-directory-build/ecosoc-snapshot.xlsx     (6,674 orgs: General 132,
      Special 5,575, Roster 967 — full ECOSOC expansion)
  ~/workspace/ngo-directory-build/ny-geneva-snapshot.xlsx (463 orgs, NY + Geneva)

Builds (idempotent):
  ngo-directory/index.html
  ngo-directory/ecosoc-global.html
  ngo-directory/new-york-geneva.html

Wires (idempotent): site-wide "More Departments" nav, homepage dept-tile rail,
sitemap.xml, data/search-index.json, data/luxury/ai-index.json.

Page template (masthead/nav/hero/footer/CSS) is lifted from
diplomatic-corps/bangkok.html so the department matches the site byte-for-byte.

When the ECOSOC expansion lands: re-copy the live workbooks over the snapshots
and rerun this script. (Expansion landed 2026-10-30: full 6,674-org dataset.)
"""
import html as htmlmod
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
SNAP = Path.home() / "workspace" / "ngo-directory-build"
ECOSOC_XLSX = SNAP / "ecosoc-snapshot.xlsx"
NYGVA_XLSX = SNAP / "ny-geneva-snapshot.xlsx"
TEMPLATE_PAGE = ROOT / "diplomatic-corps" / "bangkok.html"
OUTDIR = ROOT / "ngo-directory"
SITE = "https://likeleaders.com"
TODAY = "2026-10-30"

NGO_NAV_LINK = '<a href="/ngo-directory/index.html">NGO Directory</a>'
NAV_ANCHOR_6 = '<a href="/sections/music.html">Music</a>\n      <a href="/sections/science.html">Science</a>'
NAV_REPL_6 = ('<a href="/sections/music.html">Music</a>\n'
              '      ' + NGO_NAV_LINK + '\n'
              '      <a href="/sections/science.html">Science</a>')
NAV_ANCHOR_2 = '  <a href="/sections/music.html">Music</a>\n  <a href="/sections/film.html">Film</a>'
NAV_REPL_2 = ('  <a href="/sections/music.html">Music</a>\n'
              '  ' + NGO_NAV_LINK + '\n'
              '  <a href="/sections/film.html">Film</a>')
TILE_ANCHOR = ('<span class="label">Department</span><strong>Music</strong></a>'
               '<a class="dept-tile" href="/sections/science.html">')
TILE_REPL = ('<span class="label">Department</span><strong>Music</strong></a>'
             '<a class="dept-tile" href="/ngo-directory/index.html">'
             '<span class="label">Department</span><strong>NGO Directory</strong></a>'
             '<a class="dept-tile" href="/sections/science.html">')


def esc(v):
    if v is None:
        return ""
    return htmlmod.escape(str(v).strip(), quote=True)


def norm_url(v):
    """Normalize a website value to an href; return (href, display).

    Bracketed editorial notes (e.g. '[none published - profile shows ZAP
    placeholder]') are NOT links: display as published, no href.
    """
    v = (str(v).strip() if v else "")
    if not v:
        return "", ""
    if v.startswith("["):
        return "", v
    if re.search(r"[\s\\]", v):
        # not a single clean URL (multiple URLs, stray spaces or backslashes
        # in the published value): display as published, never fabricate a link
        return "", v
    if re.match(r"^https?://", v, re.I):
        return v, v
    return "https://" + v, v


def email_href(v):
    """Extract a bare email for the mailto href; display the full published value."""
    v = (str(v).strip() if v else "")
    if not v:
        return "", ""
    m = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", v)
    return ("mailto:" + m.group(0)) if m else "", v


def load_sheet(xlsx, sheet):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c else "" for c in rows[0]]
    data = []
    for r in rows[1:]:
        d = dict(zip(header, r))
        if not (d.get(header[0]) or d.get(header[1])):
            continue
        data.append(d)
    return header, data


# ---------------------------------------------------------------- template
_template = TEMPLATE_PAGE.read_text(encoding="utf-8")
_TOP, _FOOTER_TAIL = _template.split("</nav>", 1)
FOOTER = "<footer" + _FOOTER_TAIL.split("<footer", 1)[1]
HEAD_STYLE = _TOP.split("</head>")[0] + "</head>"
MAST_NAV = _TOP.split("</head>", 1)[1] + "</nav>"
# add the NGO Directory link to the template's own nav (alphabetical: Music -> NGO -> Science)
MAST_NAV = MAST_NAV.replace(NAV_ANCHOR_6, NAV_REPL_6)

JSONLD_RE = re.compile(r'<script type="application/ld\+json">.*?</script>', re.S)


def make_head(title, description, canonical_path, jsonld):
    head = HEAD_STYLE
    head = head.replace(
        '<link rel="canonical" href="https://likeleaders.com/diplomatic-corps/bangkok.html"/>',
        f'<link rel="canonical" href="{SITE}{canonical_path}"/>')
    head = re.sub(r"<title>.*?</title>", f"<title>{esc(title)}</title>", head, count=1)
    head = re.sub(r'<meta name="description" content=".*?"/>',
                  f'<meta name="description" content="{esc(description)}"/>', head, count=1)
    head = re.sub(r'<meta property="og:title" content=".*?"/>',
                  f'<meta property="og:title" content="{esc(title)}"/>', head, count=1)
    head = re.sub(r'<meta property="og:description" content=".*?"/>',
                  f'<meta property="og:description" content="{esc(description)}"/>', head, count=1)
    head = re.sub(r'<meta property="og:url" content=".*?"/>',
                  f'<meta property="og:url" content="{SITE}{canonical_path}"/>', head, count=1)
    head = JSONLD_RE.sub(
        lambda m: '<script type="application/ld+json">\n' + jsonld + '\n</script>',
        head, count=1)
    return head


def page_shell(title, description, canonical_path, jsonld, main_html):
    return (make_head(title, description, canonical_path, jsonld)
            + MAST_NAV + "\n<main class=\"wrap\">\n" + main_html + "\n</main>\n\n" + FOOTER)


def breadcrumbs(*crumbs):
    els = []
    for i, (name, path) in enumerate(crumbs, 1):
        els.append({"@type": "ListItem", "position": i, "name": name,
                    "item": SITE + path})
    return {"@type": "BreadcrumbList", "itemListElement": els}


def cell_link(href, text):
    if not href or not text:
        return esc(text)
    return f'<a href="{esc(href)}">{esc(text)}</a>'


# ---------------------------------------------------------------- ECOSOC page
# General status: full public detail (23-col sheet; page shows the substantive fields)
ECOSOC_COLS_GENERAL = [
    ("Organization name", "Organization"),
    ("Acronym", "Acronym"),
    ("ECOSOC status", "ECOSOC status"),
    ("Year accredited", "Year accredited"),
    ("Telephone (organization)", "Telephone"),
    ("Email (organization)", "Email"),
    ("Website", "Website"),
    ("Headquarters address (as published)", "Headquarters address"),
    ("Country (as published in address)", "Country"),
    ("Mission statement (as published)", "Mission statement"),
    ("Fields of activity / focus areas", "Fields of activity"),
    ("iCSO profile link", "iCSO profile"),
    ("Row notes", "Notes"),
]
# Special & Roster: directory-level detail (12-col sheets). The ECOSOC status
# column is redundant (the section header carries it), so it is not displayed.
ECOSOC_COLS_DIR = [
    ("Organization name", "Organization"),
    ("Year accredited", "Year accredited"),
    ("Country (as published in address)", "Country"),
    ("Headquarters address (as published)", "Headquarters address"),
    ("Telephone (organization)", "Telephone"),
    ("Email (organization)", "Email"),
    ("Website", "Website"),
    ("Organization type", "Organization type"),
    ("iCSO profile link", "iCSO profile"),
    ("Row notes", "Notes"),
]

ECOSOC_SCOPE_NOTE = (
    "Scope: the full UN DESA iCSO live database of non-governmental organizations in "
    "consultative status with the United Nations Economic and Social Council, as of "
    "30 October 2026 \u2014 6,674 organizations: General consultative status (132), "
    "Special consultative status (5,575), Roster (967). 499 records carry an extraction "
    "flag where a profile page could not be re-fetched; the registry\u2019s name, status "
    "and year are shown with the flag, never padded. A blank field means the detail is "
    "not publicly documented; contact information is organization-level only, exactly as "
    "published by the organization."
)
ECOSOC_SOURCE = (
    "Source: UN DESA NGO Branch, Integrated Civil Society Organizations (iCSO) System "
    "(esango.un.org), public listing retrieved 30 October 2026; cross-reference: the "
    "Secretary-General\u2019s official list E/2022/INF/5 (as of 31 December 2022). "
    "Directory compiled 30 October 2026. Verified facts only."
)


def ecosoc_row(r, cols):
    tds = []
    for key, _ in cols:
        v = r.get(key)
        if key == "Website":
            href, disp = norm_url(v)
            tds.append(f"<td>{cell_link(href, disp)}</td>")
        elif key == "Email (organization)":
            href, disp = email_href(v)
            tds.append(f"<td>{cell_link(href, disp)}</td>")
        elif key == "iCSO profile link":
            tds.append(f"<td>{cell_link(v, 'iCSO profile') if v else ''}</td>")
        else:
            tds.append(f"<td>{esc(v)}</td>")
    return "<tr>" + "".join(tds) + "</tr>"


def ecosoc_table(rows, cols, section_title):
    thead = "<tr>" + "".join(f"<th>{esc(label)}</th>" for _, label in cols) + "</tr>"
    body = "\n".join(ecosoc_row(r, cols) for r in rows)
    return (f'<h2 class="section-title">{section_title}</h2>\n'
            f'<div class="dna-tablewrap"><table class="dna-table">\n<thead>{thead}</thead>\n<tbody>\n'
            + body + "\n</tbody>\n</table></div>")


def ecosoc_items(rows, list_name):
    items = []
    for pos, r in enumerate(rows, 1):
        name = str(r.get("Organization name") or "").strip()
        href, _ = norm_url(r.get("Website"))
        item = {"@type": "NGO", "name": name}
        if href:
            item["url"] = href
        items.append({"@type": "ListItem", "position": pos, "item": item})
    return {"@type": "ItemList", "name": list_name,
            "url": SITE + "/ngo-directory/ecosoc-global.html",
            "numberOfItems": len(rows), "itemListElement": items}


def build_ecosoc():
    _, general = load_sheet(ECOSOC_XLSX, "General-Status NGOs")
    _, special = load_sheet(ECOSOC_XLSX, "Special-Status NGOs")
    _, roster = load_sheet(ECOSOC_XLSX, "Roster NGOs")
    assert len(general) == 132, f"expected 132 General rows, got {len(general)}"
    assert len(special) == 5575, f"expected 5575 Special rows, got {len(special)}"
    assert len(roster) == 967, f"expected 967 Roster rows, got {len(roster)}"
    n_total = len(general) + len(special) + len(roster)

    tables = (
        ecosoc_table(general, ECOSOC_COLS_GENERAL,
                     f"General consultative status \u2014 {len(general)} organizations") + "\n"
        + ecosoc_table(special, ECOSOC_COLS_DIR,
                       f"Special consultative status \u2014 {len(special)} organizations") + "\n"
        + ecosoc_table(roster, ECOSOC_COLS_DIR,
                       f"Roster \u2014 {len(roster)} organizations")
    )
    main = (
        '<p class="crumbs"><a href="/index.html">Home</a> \u00b7 '
        '<a href="/ngo-directory/index.html">NGO Directory</a> \u00b7 ECOSOC Global</p>\n'
        '<header class="dept-opener">\n'
        '  <p class="dept-kicker">NGO Directory</p>\n'
        '  <h1 class="dept-title">ECOSOC Consultative-Status NGOs \u2014 Global</h1>\n'
        '  <p class="dept-deck">Every non-governmental organization in consultative status '
        'with the United Nations Economic and Social Council \u2014 General, Special and '
        'Roster \u2014 from the UN DESA iCSO live database, listed alphabetically within each status.</p>\n'
        f'  <p class="dept-count">{n_total:,} organizations \u00b7 General {len(general)} '
        f'\u00b7 Special {len(special):,} \u00b7 Roster {len(roster)}</p>\n'
        '</header>\n'
        f'<p class="dna-note">{esc(ECOSOC_SCOPE_NOTE)}</p>\n'
        f'<p class="dna-note">{esc(ECOSOC_SOURCE)}</p>\n'
        + tables + '\n'
        '<p class="dna-note">LIKE NGO Directory. Verified facts only.</p>'
    )
    jsonld = json.dumps({
        "@context": "https://schema.org",
        "@graph": [
            breadcrumbs(("Home", "/index.html"),
                        ("NGO Directory", "/ngo-directory/index.html"),
                        ("ECOSOC Global", "/ngo-directory/ecosoc-global.html")),
            ecosoc_items(general,
                         "ECOSOC Consultative-Status NGOs \u2014 Global: General status "
                         "\u2014 LIKE Magazine"),
            ecosoc_items(special,
                         "ECOSOC Consultative-Status NGOs \u2014 Global: Special status "
                         "\u2014 LIKE Magazine"),
            ecosoc_items(roster,
                         "ECOSOC Consultative-Status NGOs \u2014 Global: Roster "
                         "\u2014 LIKE Magazine"),
        ]}, ensure_ascii=False, indent=2)
    title = "ECOSOC Consultative-Status NGOs \u2014 Global \u00b7 LIKE Magazine"
    desc = ("6,674 non-governmental organizations in consultative status with the UN Economic "
            "and Social Council \u2014 General 132, Special 5,575, Roster 967 \u2014 from the "
            "UN DESA iCSO live database. Verified facts only.")
    return (page_shell(title, desc, "/ngo-directory/ecosoc-global.html", jsonld, main),
            {"general": len(general), "special": len(special),
             "roster": len(roster), "total": n_total})


# ---------------------------------------------------------------- NY/Geneva page
NYGVA_COLS = [
    ("Organization name", "Organization"),
    ("Acronym", "Acronym"),
    ("Telephone (organization)", "Telephone"),
    ("Email (organization)", "Email"),
    ("Website", "Website"),
    ("Physical address (hub city)", "Address (hub city)"),
    ("Country of headquarters", "HQ country"),
    ("Mission / focus areas", "Mission / focus"),
    ("UN affiliation status", "UN affiliation"),
]

NYGVA_SOURCE = (
    "Sources: UN DESA NGO Branch iCSO System (esango.un.org), public search; "
    "UN Department of Global Communications associated-CSO list, February 2026; "
    "R\u00e9publique et canton de Gen\u00e8ve, geneve-int.ch Who\u2019s who directory. "
    "Retrieved 30 October 2026. Directory compiled 30 October 2026. Verified facts only."
)
NYGVA_NOTE = (
    "New York entries hold a current headquarters address verifiably in New York City; "
    "Geneva entries are listed as non-governmental organizations in the canton\u2019s official "
    "Who\u2019s who. A blank field means the detail is not publicly documented; contact "
    "information is organization-level only, exactly as published."
)


def nygva_table(rows):
    thead = "<tr>" + "".join(f"<th>{esc(label)}</th>" for _, label in NYGVA_COLS) + "</tr>"
    tbody_rows = []
    for r in rows:
        tds = []
        for key, _ in NYGVA_COLS:
            v = r.get(key)
            if key == "Website":
                href, disp = norm_url(v)
                tds.append(f"<td>{cell_link(href, disp)}</td>")
            elif key == "Email (organization)":
                href, disp = email_href(v)
                tds.append(f"<td>{cell_link(href, disp)}</td>")
            else:
                tds.append(f"<td>{esc(v)}</td>")
        tbody_rows.append("<tr>" + "".join(tds) + "</tr>")
    return (f'<div class="dna-tablewrap"><table class="dna-table">\n<thead>{thead}</thead>\n<tbody>\n'
            + "\n".join(tbody_rows) + "\n</tbody>\n</table></div>")


def build_nygva():
    _, rows = load_sheet(NYGVA_XLSX, "Organizations")
    assert len(rows) == 463, f"expected 463 NY/Geneva rows, got {len(rows)}"
    ny = [r for r in rows if (r.get("Hub") or "").strip() == "New York"]
    gva = [r for r in rows if (r.get("Hub") or "").strip() == "Geneva"]
    assert len(ny) == 330 and len(gva) == 133, f"hub split off: {len(ny)} NY, {len(gva)} Geneva"

    main = (
        '<p class="crumbs"><a href="/index.html">Home</a> \u00b7 '
        '<a href="/ngo-directory/index.html">NGO Directory</a> \u00b7 New York &amp; Geneva</p>\n'
        '<header class="dept-opener">\n'
        '  <p class="dept-kicker">NGO Directory</p>\n'
        '  <h1 class="dept-title">NGOs in New York &amp; Geneva</h1>\n'
        '  <p class="dept-deck">Non-governmental organizations with a verified presence in the '
        'two great NGO hubs \u2014 New York and Geneva.</p>\n'
        f'  <p class="dept-count">{len(ny)} New York \u00b7 {len(gva)} Geneva \u00b7 {len(rows)} organizations</p>\n'
        '</header>\n'
        f'<p class="dna-note">{esc(NYGVA_NOTE)}</p>\n'
        f'<p class="dna-note">{esc(NYGVA_SOURCE)}</p>\n'
        f'<h2 class="section-title">New York \u2014 {len(ny)} organizations</h2>\n' + nygva_table(ny) + '\n'
        f'<h2 class="section-title">Geneva \u2014 {len(gva)} organizations</h2>\n' + nygva_table(gva) + '\n'
        '<p class="dna-note">LIKE NGO Directory. Verified facts only.</p>'
    )
    items = []
    for pos, r in enumerate(rows, 1):
        name = str(r.get("Organization name") or "").strip()
        href, _ = norm_url(r.get("Website"))
        item = {"@type": "NGO", "name": name}
        if href:
            item["url"] = href
        items.append({"@type": "ListItem", "position": pos, "item": item})
    jsonld = json.dumps({
        "@context": "https://schema.org",
        "@graph": [
            breadcrumbs(("Home", "/index.html"),
                        ("NGO Directory", "/ngo-directory/index.html"),
                        ("New York & Geneva", "/ngo-directory/new-york-geneva.html")),
            {"@type": "ItemList",
             "name": "NGOs in New York & Geneva \u2014 LIKE Magazine",
             "url": SITE + "/ngo-directory/new-york-geneva.html",
             "numberOfItems": len(rows),
             "itemListElement": items},
        ]}, ensure_ascii=False, indent=2)
    title = "NGOs in New York & Geneva \u00b7 LIKE Magazine"
    desc = ("463 non-governmental organizations with a verified presence in New York (330) "
            "and Geneva (133). Verified facts only.")
    return page_shell(title, desc, "/ngo-directory/new-york-geneva.html", jsonld, main), len(rows)


# ---------------------------------------------------------------- hub index page
def build_index():
    cards = [
        ("/ngo-directory/ecosoc-global.html", "NGO Directory",
         "ECOSOC Consultative-Status NGOs \u2014 Global",
         "6,674 organizations \u00b7 General 132 \u00b7 Special 5,575 \u00b7 Roster 967"),
        ("/ngo-directory/new-york-geneva.html", "NGO Directory",
         "NGOs in New York &amp; Geneva",
         "330 New York \u00b7 133 Geneva \u00b7 463 organizations"),
    ]
    grid = "\n".join(
        f'<a class="edition" href="{href}"><span class="ed-num">{kicker}</span>'
        f'<strong>{name}</strong><em>{sub}</em></a>'
        for href, kicker, name, sub in cards)
    main = (
        '<p class="crumbs"><a href="/index.html">Home</a> \u00b7 NGO Directory</p>\n'
        '<header class="dept-opener">\n'
        '  <p class="dept-kicker">LIKE Magazine \u00b7 Directory Department</p>\n'
        '  <h1 class="dept-title">NGO Directory</h1>\n'
        '  <p class="dept-deck">The organizations that carry civil society: non-governmental '
        'organizations in consultative status with the United Nations Economic and Social Council, '
        'and the NGO communities of New York and Geneva \u2014 gathered from official sources, '
        'checked, and kept current.</p>\n'
        '  <p class="dept-count">2 lists \u00b7 7,137 organizations</p>\n'
        '</header>\n'
        '<h2 class="section-title">The lists</h2>\n'
        f'<div class="edition-grid">\n{grid}\n</div>\n'
        '<p class="dna-note">Contact information is organization-level only, exactly as published; '
        'a blank field means the detail is not publicly documented. More lists follow as each is validated.</p>\n'
        '<p class="dna-note">LIKE NGO Directory. Verified facts only.</p>'
    )
    jsonld = json.dumps({
        "@context": "https://schema.org",
        "@graph": [
            breadcrumbs(("Home", "/index.html"), ("NGO Directory", "/ngo-directory/index.html")),
            {"@type": "CollectionPage",
             "name": "NGO Directory \u2014 LIKE Magazine",
             "url": SITE + "/ngo-directory/index.html",
             "description": "Non-governmental organizations in UN ECOSOC consultative status and "
                            "the NGO hubs of New York and Geneva, from official sources.",
             "hasPart": [
                 {"@type": "WebPage", "name": "ECOSOC Consultative-Status NGOs \u2014 Global",
                  "url": SITE + "/ngo-directory/ecosoc-global.html"},
                 {"@type": "WebPage", "name": "NGOs in New York & Geneva",
                  "url": SITE + "/ngo-directory/new-york-geneva.html"},
             ]},
        ]}, ensure_ascii=False, indent=2)
    title = "NGO Directory \u2014 Worldwide Directories \u00b7 LIKE Magazine"
    desc = ("LIKE NGO Directory: non-governmental organizations with UN ECOSOC consultative "
            "status and the NGO hubs of New York and Geneva \u2014 7,137 organizations. "
            "Verified facts only.")
    return page_shell(title, desc, "/ngo-directory/index.html", jsonld, main)


# ---------------------------------------------------------------- wiring
def patch_nav():
    changed = 0
    for f in ROOT.rglob("*.html"):
        if "ngo-directory" in str(f):
            continue
        h = f.read_text(encoding="utf-8", errors="replace")
        if "/ngo-directory/index.html" in h:
            continue
        if NAV_ANCHOR_6 in h:
            f.write_text(h.replace(NAV_ANCHOR_6, NAV_REPL_6), encoding="utf-8")
            changed += 1
        elif NAV_ANCHOR_2 in h:
            f.write_text(h.replace(NAV_ANCHOR_2, NAV_REPL_2), encoding="utf-8")
            changed += 1
    return changed


def patch_homepage_tile():
    p = ROOT / "index.html"
    h = p.read_text(encoding="utf-8")
    tile_sig = 'href="/ngo-directory/index.html"><span class="label">Department</span>'
    if tile_sig in h or TILE_ANCHOR not in h:
        return 0
    p.write_text(h.replace(TILE_ANCHOR, TILE_REPL), encoding="utf-8")
    return 1


def patch_sitemap():
    p = ROOT / "sitemap.xml"
    xml = p.read_text(encoding="utf-8")
    added = 0
    entries = [
        ("/ngo-directory/index.html", "0.9"),
        ("/ngo-directory/ecosoc-global.html", "0.8"),
        ("/ngo-directory/new-york-geneva.html", "0.8"),
    ]
    anchor = "    <loc>https://likeleaders.com/sections/music.html</loc>"
    for path, prio in entries:
        if path in xml:
            continue
        entry = (f'  <url>\n    <loc>{SITE}{path}</loc>\n'
                 f'    <lastmod>{TODAY}</lastmod>\n    <changefreq>monthly</changefreq>\n'
                 f'    <priority>{prio}</priority>\n  </url>\n')
        i = xml.find(anchor)
        j = xml.find("</url>", i) + len("</url>")
        xml = xml[:j] + "\n" + entry.rstrip("\n") + xml[j:]
        added += 1
    if added:
        p.write_text(xml, encoding="utf-8")
    return added


def patch_search_index():
    p = ROOT / "data" / "search-index.json"
    idx = json.loads(p.read_text(encoding="utf-8"))
    entries = [
        ("/ngo-directory/index.html",
         "NGO Directory \u2014 Worldwide Directories \u00b7 LIKE Magazine",
         "LIKE NGO Directory: non-governmental organizations with UN ECOSOC consultative "
         "status and the NGO hubs of New York and Geneva \u2014 7,137 organizations. Verified facts only.",
         "NGO Directory"),
        ("/ngo-directory/ecosoc-global.html",
         "ECOSOC Consultative-Status NGOs \u2014 Global \u00b7 LIKE Magazine",
         "6,674 UN ECOSOC consultative-status NGOs \u2014 General 132, Special 5,575, Roster 967 "
         "\u2014 from the UN DESA iCSO live database. Verified facts only.",
         "NGO Directory"),
        ("/ngo-directory/new-york-geneva.html",
         "NGOs in New York & Geneva \u00b7 LIKE Magazine",
         "463 non-governmental organizations with a verified presence in New York (330) "
         "and Geneva (133). Verified facts only.",
         "NGO Directory"),
    ]
    changed = 0
    for url, title, desc, kicker in entries:
        hit = next((e for e in idx if e.get("url") == url), None)
        if hit is None:
            idx.append({"url": url, "title": title, "description": desc, "kicker": kicker})
            changed += 1
        elif hit.get("title") != title or hit.get("description") != desc:
            hit["title"] = title
            hit["description"] = desc
            hit["kicker"] = kicker
            changed += 1
    if changed:
        p.write_text(json.dumps(idx, ensure_ascii=False, indent=1), encoding="utf-8")
    return changed


def patch_ai_index():
    p = ROOT / "data" / "luxury" / "ai-index.json"
    d = json.loads(p.read_text(encoding="utf-8"))
    ds = d.get("dataset", [])
    entries = [
        ("UN ECOSOC Consultative-Status NGOs \u2014 Global",
         "Non-governmental organizations in consultative status with the UN Economic "
         "and Social Council \u2014 General 132, Special 5,575, Roster 967; 6,674 total, "
         "from the UN DESA iCSO live database.",
         "/ngo-directory/ecosoc-global.html", 6674),
        ("NGOs in New York & Geneva",
         "Non-governmental organizations with a verified presence in New York (330) and Geneva (133).",
         "/ngo-directory/new-york-geneva.html", 463),
    ]
    changed = 0
    for name, desc, path, count in entries:
        hit = next((x for x in ds if x.get("name") == name), None)
        if hit is None:
            ds.append({
                "@type": "Dataset",
                "name": name,
                "description": desc,
                "url": SITE + path,
                "creator": {"@id": "https://likeleaders.com/#organization"},
                "isPartOf": {"@type": "CreativeWork", "name": "LIKE\u00ae Magazine NGO Directory",
                             "url": f"{SITE}/ngo-directory/index.html"},
                "distribution": {"@type": "DataDownload", "encodingFormat": "text/html",
                                 "contentUrl": SITE + path},
                "variableMeasured": "listing",
                "count": count,
            })
            changed += 1
        elif hit.get("description") != desc or hit.get("count") != count:
            hit["description"] = desc
            hit["count"] = count
            changed += 1
    if changed:
        d["dataset"] = ds
        d["dateModified"] = TODAY
        p.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")
    return changed


# ---------------------------------------------------------------- validation
PLACEHOLDER_RE = re.compile(
    r"lorem ipsum|TODO|FIXME|XXX|placeholder|Your Name|example\.com", re.I)


def template_text(h):
    """Page text minus data: JSON-LD, table cells, and bracketed workbook
    editorial notes (e.g. ZAP placeholders) are legitimate published content,
    not template placeholders. What remains is the template's own wording,
    which is what the placeholder check guards."""
    h = JSONLD_RE.sub("", h)
    h = re.sub(r"<td>.*?</td>", "", h, flags=re.S)
    return re.sub(r"\[[^\]]*ZAP[^\]]*\]", "", h)


def validate_page(path, expected_items=None, must_contain=(), must_not_contain=()):
    h = path.read_text(encoding="utf-8")
    errs = []
    m = JSONLD_RE.search(h)
    if not m:
        errs.append("no JSON-LD")
    else:
        try:
            data = json.loads(m.group(0).split(">", 1)[1].rsplit("<", 1)[0])
            graph = data.get("@graph", [])
            il = next((g for g in graph if g.get("@type") == "ItemList"), None)
            if expected_items is not None:
                if not il:
                    errs.append("no ItemList in JSON-LD")
                elif il.get("numberOfItems") != expected_items:
                    errs.append(f"numberOfItems {il.get('numberOfItems')} != {expected_items}")
                elif len(il.get("itemListElement", [])) != expected_items:
                    errs.append("itemListElement count mismatch")
        except Exception as e:
            errs.append(f"JSON-LD parse error: {e}")
    for s in must_contain:
        if s not in h:
            errs.append(f"missing required text: {s[:60]}")
    for s in must_not_contain:
        if s in h:
            errs.append(f"forbidden text still present: {s[:60]}")
    if PLACEHOLDER_RE.search(template_text(h)):
        errs.append("placeholder text found")
    if "ngo-directory" not in h:
        errs.append("self link missing")
    return errs


def validate_ecosoc_page(path):
    """ECOSOC page carries three ItemLists (General 132 / Special 5575 / Roster 967)."""
    h = path.read_text(encoding="utf-8")
    errs = []
    m = JSONLD_RE.search(h)
    if not m:
        return ["no JSON-LD"]
    try:
        data = json.loads(m.group(0).split(">", 1)[1].rsplit("<", 1)[0])
    except Exception as e:
        return [f"JSON-LD parse error: {e}"]
    graph = data.get("@graph", [])
    ils = [g for g in graph if g.get("@type") == "ItemList"]
    expect = {"General status": 132, "Special status": 5575, "Roster": 967}
    if len(ils) != 3:
        errs.append(f"expected 3 ItemLists, found {len(ils)}")
    for il in ils:
        name = il.get("name", "")
        key = next((k for k in expect if k in name), None)
        if key is None:
            errs.append(f"unexpected ItemList name: {name[:60]}")
            continue
        if il.get("numberOfItems") != expect[key]:
            errs.append(f"{key}: numberOfItems {il.get('numberOfItems')} != {expect[key]}")
        if len(il.get("itemListElement", [])) != expect[key]:
            errs.append(f"{key}: itemListElement count mismatch")
    for s in ["6,674", "General 132", "Special 5,575", "Roster 967", "extraction"]:
        if s not in h:
            errs.append(f"missing required text: {s[:60]}")
    if "expansion underway" in h:
        errs.append("forbidden text still present: expansion underway")
    if PLACEHOLDER_RE.search(template_text(h)):
        errs.append("placeholder text found")
    return errs


def main():
    OUTDIR.mkdir(exist_ok=True)
    index_html = build_index()
    ecosoc_html, ecosoc_counts = build_ecosoc()
    nygva_html, n_nygva = build_nygva()
    (OUTDIR / "index.html").write_text(index_html, encoding="utf-8")
    (OUTDIR / "ecosoc-global.html").write_text(ecosoc_html, encoding="utf-8")
    (OUTDIR / "new-york-geneva.html").write_text(nygva_html, encoding="utf-8")
    print(f"built ngo-directory/index.html")
    print(f"built ngo-directory/ecosoc-global.html: {ecosoc_counts}")
    print(f"built ngo-directory/new-york-geneva.html: {n_nygva} orgs")
    print("nav patched:", patch_nav())
    print("homepage tile:", patch_homepage_tile())
    print("sitemap entries added:", patch_sitemap())
    print("search index updated:", patch_search_index())
    print("ai-index updated:", patch_ai_index())

    errs = []
    errs += [("index.html", e) for e in validate_page(
        OUTDIR / "index.html", must_contain=["NGO Directory", "7,137 organizations"],
        must_not_contain=["expansion underway", "488 organizations"])]
    errs += [("ecosoc-global.html", e) for e in validate_ecosoc_page(OUTDIR / "ecosoc-global.html")]
    errs += [("new-york-geneva.html", e) for e in validate_page(
        OUTDIR / "new-york-geneva.html", expected_items=463,
        must_contain=["330", "133"])]
    # every internal link resolves
    for pg in ["index.html", "ecosoc-global.html", "new-york-geneva.html"]:
        h = (OUTDIR / pg).read_text(encoding="utf-8")
        for link in sorted(set(re.findall(r'href="(/[^"]+)"', h))):
            target = ROOT / link.lstrip("/")
            if not target.exists():
                errs.append((pg, f"broken internal link: {link}"))
    if errs:
        print("VALIDATION ERRORS:")
        for pg, e in errs:
            print(" ", pg, "-", e)
        sys.exit(1)
    print("validation: all pass")


if __name__ == "__main__":
    main()
